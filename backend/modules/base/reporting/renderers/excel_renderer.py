"""Excel (.xlsx) report renderer with corporate styling, zebra striping, and auto column widths."""

import io
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from modules.base.reporting.registry import ReportDataResult
from modules.base.reporting.models import ReportTemplate


class ExcelReportRenderer:
    """Renders report data as styled OpenPyXL workbook."""

    @classmethod
    def render(
        cls,
        data: ReportDataResult,
        template: Optional[ReportTemplate] = None,
    ) -> bytes:
        """Render report into styled .xlsx binary bytes."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (data.report_name[:28] if len(data.report_name) > 28 else data.report_name).replace("/", "-")

        # Color theme
        primary_hex = (template.primary_color if template and template.primary_color else "#1E3A8A").lstrip("#")
        header_fill = PatternFill(start_color=primary_hex, end_color=primary_hex, fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10)
        font_total = Font(name="Calibri", size=10, bold=True)
        font_title = Font(name="Calibri", size=16, bold=True, color=primary_hex)
        font_subtitle = Font(name="Calibri", size=9, italic=True, color="64748B")

        thin_side = Side(border_style="thin", color="CBD5E1")
        double_side = Side(border_style="double", color="475569")
        cell_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        total_border = Border(top=thin_side, bottom=double_side, left=thin_side, right=thin_side)

        current_row = 1

        # 1. Company Branding & Report Title
        company_name = (data.company_info or {}).get("name", "Sovereign Platform")
        ws.cell(row=current_row, column=1, value=company_name).font = font_title
        current_row += 1

        ws.cell(
            row=current_row,
            column=1,
            value=f"{data.report_name} | Generated: {data.generated_at[:19]}",
        ).font = font_subtitle
        current_row += 2

        # 2. Table Column Headers
        header_row_idx = current_row
        for col_idx, col in enumerate(data.columns, start=1):
            cell = ws.cell(row=header_row_idx, column=col_idx, value=col["title"])
            cell.font = font_header
            cell.fill = header_fill
            cell.border = cell_border
            align_h = "right" if col["type"] in ("integer", "float") else "left"
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=True)

        ws.row_dimensions[header_row_idx].height = 24
        current_row += 1

        # 3. Data Rows
        col_names = [col["name"] for col in data.columns]
        for row_idx, row_data in enumerate(data.rows):
            is_even = row_idx % 2 == 1
            for col_idx, col in enumerate(data.columns, start=1):
                c_name = col["name"]
                c_type = col["type"]
                val = row_data.get(c_name)

                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = font_data
                cell.border = cell_border
                if is_even:
                    cell.fill = zebra_fill

                align_h = "right" if c_type in ("integer", "float") else "left"
                cell.alignment = Alignment(horizontal=align_h, vertical="center")

                # Number formatting
                if c_type == "float" and isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
                elif c_type == "integer" and isinstance(val, int):
                    cell.number_format = "#,##0"

            ws.row_dimensions[current_row].height = 20
            current_row += 1

        # 4. Aggregates / Totals Row
        if data.aggregates:
            for col_idx, col in enumerate(data.columns, start=1):
                c_name = col["name"]
                c_type = col["type"]
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = font_total
                cell.border = total_border
                cell.fill = total_fill

                if c_name in data.aggregates:
                    agg_val = data.aggregates[c_name]
                    cell.value = agg_val
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if c_type == "float":
                        cell.number_format = "#,##0.00"
                    elif c_type == "integer":
                        cell.number_format = "#,##0"
                elif col_idx == 1:
                    cell.value = "TOTAL"
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.value = ""

            ws.row_dimensions[current_row].height = 22

        # 5. Auto-adjust column widths based on cell text lengths
        for col_idx, col in enumerate(data.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(col["title"])
            for r in range(header_row_idx, current_row + 1):
                v = ws.cell(row=r, column=col_idx).value
                if v is not None:
                    max_len = max(max_len, len(str(v)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
